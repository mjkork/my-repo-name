from io import BytesIO, StringIO

from django.contrib import messages
from django.core.management import call_command
from django.http import HttpRequest, HttpResponse, HttpResponseNotAllowed
from django.shortcuts import redirect, render
from django.utils import timezone
from django.views import View
from openpyxl import Workbook
from openpyxl.styles import Font

from preferences.forms import UserPreferencesForm
from preferences.models import UserPreferences
from sessions.models import Session


class MySettingsView(View):
    """Settings page; hosts backup management and future config features."""

    template_name = "preferences/mysettings.html"

    def get(self, request: HttpRequest) -> HttpResponse:
        prefs = UserPreferences.load()
        form = UserPreferencesForm(instance=prefs)
        return render(request, self.template_name, {"preferences_form": form})


class UpdatePreferencesView(View):
    """POST-only: save UserPreferences and redirect back to /mysettings/."""

    def post(self, request: HttpRequest) -> HttpResponse:
        prefs = UserPreferences.load()
        form = UserPreferencesForm(request.POST, instance=prefs)
        if form.is_valid():
            form.save()
            messages.success(request, "Page size updated.")
        else:
            messages.error(request, "Invalid value — preferences not saved.")
        return redirect("preferences:mysettings")

    def get(self, request: HttpRequest) -> HttpResponse:
        return HttpResponseNotAllowed(["POST"])


class BackupDownloadView(View):
    """Export all user data as a dated JSON file via dumpdata."""

    def get(self, request: HttpRequest) -> HttpResponse:
        buf = StringIO()
        call_command(
            "dumpdata",
            exclude=["contenttypes", "auth.permission", "sessions"],
            natural_foreign=True,
            natural_primary=True,
            indent=2,
            stdout=buf,
        )
        filename = f"myshots-backup-{timezone.localdate()}.json"
        response = HttpResponse(buf.getvalue(), content_type="application/json")
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response


# (column header, width, callable extracting the cell value from a Session)
EXPORT_COLUMNS = [
    ("session_id", 10, lambda s: s.pk),
    ("date", 12, lambda s: s.date.isoformat()),
    ("name", 25, lambda s: s.name),
    ("bow_name", 20, lambda s: s.bow.name if s.bow else ""),
    ("location", 12, lambda s: s.location or ""),
    ("distance_m", 12, lambda s: s.distance_m or ""),
    ("total_arrows", 13, lambda s: s.total_arrows or ""),
    ("scoring_arrows", 15, lambda s: s.scoring_arrows or ""),
    ("target_face", 14, lambda s: s.target_face or ""),
    ("total_score", 13, lambda s: s.total_score or ""),
    ("time_of_day", 13, lambda s: s.time_of_day or ""),
    ("weather", 14, lambda s: s.weather or ""),
    ("temperature_celsius", 12, lambda s: s.temperature_celsius or ""),
    ("wind_force", 11, lambda s: s.wind_force or ""),
    (
        "wind_force_label",
        16,
        lambda s: s.get_wind_force_display() if s.wind_force else "",
    ),
    ("wind_direction", 15, lambda s: s.wind_direction or ""),
    ("nutrition", 11, lambda s: s.nutrition or ""),
    ("nutrition_label", 15, lambda s: s.get_nutrition_display() if s.nutrition else ""),
    ("sleep_hours", 12, lambda s: s.sleep_hours or ""),
    ("sleep_notes", 30, lambda s: s.sleep_notes or ""),
    ("stress", 10, lambda s: s.stress or ""),
    ("stress_label", 13, lambda s: s.get_stress_display() if s.stress else ""),
    ("fatigue", 10, lambda s: s.fatigue or ""),
    ("fatigue_label", 13, lambda s: s.get_fatigue_display() if s.fatigue else ""),
    ("physical_sensations", 35, lambda s: s.physical_sensations or ""),
    ("notes", 35, lambda s: s.notes or ""),
    ("next_focus", 30, lambda s: s.next_focus or ""),
]


class ExportDownloadView(View):
    """Export all session data as a dated .xlsx file for external analysis."""

    def get(self, request: HttpRequest) -> HttpResponse:
        sessions = Session.objects.select_related("bow").order_by("date", "pk")

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sessions"

        header_font = Font(bold=True)
        for col_idx, (header, width, _) in enumerate(EXPORT_COLUMNS, start=1):
            cell = sheet.cell(row=1, column=col_idx, value=header)
            cell.font = header_font
            sheet.column_dimensions[cell.column_letter].width = width

        for row_idx, session in enumerate(sessions, start=2):
            for col_idx, (_, _, extractor) in enumerate(EXPORT_COLUMNS, start=1):
                sheet.cell(row=row_idx, column=col_idx, value=extractor(session))

        buf = BytesIO()
        workbook.save(buf)
        filename = f"myshots-export-{timezone.localdate()}.xlsx"
        response = HttpResponse(
            buf.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response
